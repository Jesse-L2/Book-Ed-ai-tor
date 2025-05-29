"""
BookAIditor Web App - A Flask application that uses Ollama with your local AI model of choice to analyze and improve writing quality.
This application allows users to upload Word documents, processes them through the Gemma 3 model via Ollama,
and provides writing improvement suggestions and edits in the browser.
"""

import os
import json
import docx
import requests
import time
import uuid
import logging
from typing import List, Dict, Any, Tuple, Optional
from flask import Flask, render_template, request, redirect, url_for, flash, session, jsonify
from pydantic import BaseModel, ValidationError

from contexts.prompt_contexts import (
    literary_fiction_context, 
    thriller_context, 
    romance_context, 
    sci_fi_context, 
    young_adult_context, 
    fantasy_context, 
    horror_context, 
    mystery_context
)

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    handlers=[logging.StreamHandler()]
)
logger = logging.getLogger(__name__)

# Define models for validation
class ParagraphAnalysis(BaseModel):
    original: str
    issues: List[str]
    suggestions: List[str]
    improved_version: str

class AnalysisResult(BaseModel):
    overall_assessment: str
    paragraphs: List[ParagraphAnalysis]

# Initialize Flask app
app = Flask(__name__)
app.secret_key = os.urandom(24)  # Secret key handling
app.config['UPLOAD_FOLDER'] = 'uploads'
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024  # 50MB max upload size
app.config['ALLOWED_EXTENSIONS'] = {'docx', 'doc', 'pdf', 'txt'}

# Ensure upload directory exists
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

def allowed_file(filename: str) -> bool:
    """Check if file extension is allowed"""
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in app.config['ALLOWED_EXTENSIONS']

def read_word_document(file_stream):
    """Read content from file stream."""
    try:
        doc = docx.Document(file_stream)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return doc, paragraphs
    except Exception as e:
        print(f"Error reading document: {e}")
        return None, []

def chunk_paragraphs(paragraphs, chunk_size):
    # TODO: double check chunking math
    """Yield overlapping chunks (sliding window) of the given size."""
    for i in range(len(paragraphs) - chunk_size + 1):
        yield paragraphs[i:i + chunk_size]

def get_additional_context(genre):
    """Determine additional injected prompt context based on genre"""
    genre_contexts = {
        'literary_fiction': literary_fiction_context,
        'thriller': thriller_context,
        'romance': romance_context,
        'sci_fi': sci_fi_context,
        'young_adult': young_adult_context,
        'fantasy': fantasy_context,
        'horror': horror_context,
        'mystery': mystery_context
    }
    return genre_contexts.get(genre, "")

def extract_json_from_text(text: str) -> Optional[Dict[str, Any]]:
    """Safely extract JSON from a text string with repair attempts"""
    try:
        # Try to find a JSON object in the text
        json_start = text.find('{')
        json_end = text.rfind('}') + 1
        
        if json_start >= 0 and json_end > json_start:
            json_content = text[json_start:json_end]
            try:
                # First attempt: Try parsing as is
                return json.loads(json_content)
            except json.JSONDecodeError as e:
                logger.warning(f"Initial JSON parse failed: {e}. Attempting to clean up the JSON string.")
                
                # Second attempt: Remove control characters and fix common issues
                import re
                # Remove control characters
                json_content = re.sub(r'[\x00-\x1F\x7F-\x9F]', '', json_content)
                # Fix common formatting issues
                json_content = json_content.replace('\\n', '\\\\n').replace('\\"', '\\\\"')
                
                try:
                    return json.loads(json_content)
                except json.JSONDecodeError:
                    logger.warning("Second JSON parse attempt failed. Using custom parser.")
                    
                    # Third attempt: Try to extract a well-formed JSON object
                    # This is a more aggressive approach
                    try:
                        # Look for "overall_assessment" and paragraphs array
                        match = re.search(r'"overall_assessment"\s*:\s*"([^"]+)"', json_content)
                        if match:
                            overall = match.group(1)
                            
                            # Construct a minimal valid JSON
                            return {
                                "overall_assessment": overall,
                                "paragraphs": [
                                    {
                                        "original": "Paragraph text",
                                        "issues": ["Unable to parse issues from model response"],
                                        "suggestions": ["Please check the original text"],
                                        "improved_version": "Please review the original document"
                                    }
                                ]
                            }
                    except Exception:
                        logger.error("Failed to extract partial JSON content")
                        return None
        else:
            logger.warning("Could not find JSON content in response")
            return None
    except Exception as e:
        logger.error(f"Error in JSON extraction: {e}")
        logger.debug(f"Raw text: {text[:200]}...")
        return None
    
def validate_analysis_json(data: Dict[str, Any]) -> bool:
    """Validate the structure of the analysis JSON data"""
    if not isinstance(data, dict):
        return False
        
    if "overall_assessment" not in data or not isinstance(data["overall_assessment"], str):
        return False
        
    if "paragraphs" not in data or not isinstance(data["paragraphs"], list):
        return False
        
    # Check each paragraph entry
    for para in data["paragraphs"]:
        if not isinstance(para, dict):
            return False
            
        # Check for required fields
        required_fields = ["original", "issues", "suggestions", "improved_version"]
        for field in required_fields:
            if field not in para:
                return False
                
        # Validate types
        if not isinstance(para["original"], str):
            return False
        if not isinstance(para["issues"], list):
            return False
        if not isinstance(para["suggestions"], list):
            return False
        if not isinstance(para["improved_version"], str):
            return False
            
    return True

def analyze_text_with_ollama(text_chunk: str, model_name: str, api_url: str, additional_context: str = "") -> Optional[AnalysisResult]:
    """
    Send text to Ollama API for analysis and improvement suggestions.
    Uses Pydantic to validate and parse the response.
    """
    prompt = f"""
You are a professional book editor. Analyze the following passage for writing quality and marketability. Give constructive, actionable feedback that tightens language, improves clarity and rhythm, and enhances narrative logic, flow, character development, pacing, and plot structure. Preserve the author's voice and intent. Avoid providing unnecessary edits or suggestions that do not directly improve the text, it would be better to just return the original text in those cases.

Additional context: {additional_context}

Analyze ONLY the provided text chunk below. Do NOT reference previous or next chunks.

TEXT TO ANALYZE:
{text_chunk}

For each paragraph, provide:
1. A list of writing quality issues (e.g., wordiness, passive voice, unclear sentences).
2. A list of suggested improvements (e.g., clearer wording, stronger verbs, better sentence structure).
3. An improved version of the paragraph.

**RESPONSE FORMAT:**
Respond ONLY with a single valid JSON object, with no extra text or explanation. All fields must be present, even if empty.

Example:
{{
  "overall_assessment": "Brief summary of the writing quality.",
  "paragraphs": [
    {{
      "original": "The original paragraph text.",
      "issues": ["Issue 1", "Issue 2"],
      "suggestions": ["Suggestion 1", "Suggestion 2"],
      "improved_version": "The improved paragraph text."
    }}
    // ...repeat for each paragraph
  ]
}}

DO NOT include any instructions, explanations, or repeated prompt text. Output ONLY the JSON object.
"""

    try:
        # Set reasonable timeout
        response = requests.post(
            f"{api_url}/api/generate",
            json={
                "model": model_name,
                "prompt": prompt,
                "stream": False,
                "options": {
                    "temperature": 0.2,
                }
            },
            timeout=120
        )
        
        if response.status_code != 200:
            logger.error(f"API request failed with status {response.status_code}: {response.text}")
            return None
            
        response_data = response.json()
        if not response_data or 'response' not in response_data:
            logger.error("Invalid response format from Ollama API")
            return None
            
        response_text = response_data.get('response', '')
        logger.debug(f"Raw response from API: {response_text[:500]}...")
        
        # Extract JSON from response
        json_data = extract_json_from_text(response_text)
        if not json_data:
            logger.error("Failed to extract valid JSON from response")
            # Create a fallback analysis result
            return create_fallback_analysis(text_chunk)
            
        try:
            # Use Pydantic to validate the structure
            if validate_analysis_json(json_data):
                analysis_result = AnalysisResult.model_validate(json_data)
                return analysis_result
            else:
                # If validation failed, use our fallback
                logger.warning("JSON structure validation failed, using fallback analysis")
                return create_fallback_analysis(text_chunk, json_data.get("overall_assessment", ""))
        except ValidationError as ve:
            logger.error(f"Pydantic validation error: {ve}")
            # Create a fallback result with basic structure
            return create_fallback_analysis(text_chunk)
            
    except requests.exceptions.RequestException as e:
        logger.error(f"Request error during API call: {e}")
        return None
    except Exception as e:
        logger.error(f"Unexpected error during API request: {e}")
        return None

def process_document(file_stream, options: Dict[str, Any]) -> Tuple[Dict[str, Any], Optional[List[str]]]:
    """Process the document and return analysis results"""
    model_name = options.get('model', 'gemma3:1b')
    api_url = options.get('api_url', 'http://localhost:11434')
    chunk_size = int(options.get('chunk_size', 3))
    genre = options.get('genre', '')
    sample_only = options.get('sample_only', False)
    
    # Read the document
    doc, paragraphs = read_word_document(file_stream)
    if not doc or not paragraphs:
        return {"error": "Failed to read document"}, None
    
    # Get additional context based on genre
    additional_context = get_additional_context(genre)
    
    # Create chunks of paragraphs
    if len(paragraphs) < chunk_size:
        # Handle case where document has fewer paragraphs than chunk size
        chunks = [paragraphs]
    else:
        chunks = list(chunk_paragraphs(paragraphs, chunk_size))
        
    if sample_only:
        chunks = chunks[:1]
    
    # Process each chunk
    analysis_results = []
    successful_chunks = 0
    max_attempts = 3  # Maximum number of retry attempts per chunk
    
    for i, chunk in enumerate(chunks):
        chunk_text = "\n\n".join(chunk)
        
        # Try multiple times to get a valid result
        for attempt in range(max_attempts):
            try:
                logger.info(f"Processing chunk {i+1}/{len(chunks)} (attempt {attempt+1}/{max_attempts})")
                result = analyze_text_with_ollama(chunk_text, model_name, api_url, additional_context)
                if result:
                    analysis_results.append(result)
                    successful_chunks += 1
                    break  # Break out of retry loop if successful
                else:
                    logger.warning(f"Analysis for chunk {i+1} returned None on attempt {attempt+1}")
                    if attempt < max_attempts - 1:
                        logger.info(f"Retrying chunk {i+1} in 2 seconds...")
                        time.sleep(2)  # Wait before retrying
            except Exception as e:
                logger.error(f"Error processing chunk {i+1} (attempt {attempt+1}): {e}")
                if attempt < max_attempts - 1:
                    logger.info(f"Retrying chunk {i+1} in 2 seconds...")
                    time.sleep(2)  # Wait before retrying
        
        # Brief pause to avoid overwhelming the API
        if i < len(chunks) - 1:
            time.sleep(1)
    
    # Check if we got any valid results
    if not analysis_results:
        return {"error": "Failed to generate analysis results. Please try again or check if Ollama is running."}, None
    
    # Prepare results for display
    formatted_results = {
        "overall_assessment": analysis_results[0].overall_assessment if analysis_results else "No overall assessment available.",
        "paragraphs": [],
        "meta": {
            "total_chunks": len(chunks),
            "successful_chunks": successful_chunks,
            "model": model_name,
            "timestamp": time.strftime("%Y-%m-%d %H:%M:%S"),
        }
    }
    
    # Track processed paragraphs to avoid duplicates
    processed_paragraph_indices = set()
    
    for chunk_result in analysis_results:
        if not chunk_result or not hasattr(chunk_result, "paragraphs"):
            continue
        
        for para_analysis in chunk_result.paragraphs:
            # Find the matching paragraph in the original document
            # This is more robust than relying on indices
            original_text = para_analysis.original.strip()
            
            # Find best matching paragraph
            best_match_idx = -1
            best_match_score = 0
            
            for idx, para in enumerate(paragraphs):
                if idx in processed_paragraph_indices:
                    continue
                    
                # Simple similarity check - could be improved
                if original_text in para or para in original_text:
                    # Calculate a simple similarity score based on length difference
                    similarity = min(len(para), len(original_text)) / max(len(para), len(original_text))
                    if similarity > best_match_score:
                        best_match_score = similarity
                        best_match_idx = idx
            
            # If we found a match
            if best_match_idx >= 0:
                formatted_para = {
                    "index": best_match_idx + 1,
                    "original": paragraphs[best_match_idx],
                    "issues": para_analysis.issues,
                    "suggestions": para_analysis.suggestions,
                    "improved_version": para_analysis.improved_version
                }
                formatted_results["paragraphs"].append(formatted_para)
                processed_paragraph_indices.add(best_match_idx)
    
    # Sort paragraphs by index for proper ordering
    formatted_results["paragraphs"].sort(key=lambda p: p["index"])
    
    return formatted_results, paragraphs

@app.route('/')
def index():
    """Render the upload form"""
    return render_template('index.html')

@app.route('/analyze', methods=['POST'])
def analyze():
    """Handle the document upload and analysis"""
    # Check if a file was uploaded
    if 'document' not in request.files:
        flash('No file selected')
        return redirect(request.url)
    
    file = request.files['document']
    
    if file.filename == '':
        flash('No file selected')
        return redirect(request.url)
    
    if file and allowed_file(file.filename):
        # Get options from form
        options = {
            'model': request.form.get('model', 'gemma3:1b'),
            'api_url': request.form.get('api_url', 'http://localhost:11434'),
            'chunk_size': request.form.get('chunk_size', '3'),
            'genre': request.form.get('genre', ''),
            'sample_only': 'sample_only' in request.form
        }
        
        # Add focus areas if specified
        focus_areas = []
        for focus in ['grammar', 'style', 'pacing', 'characters']:
            if f'focus_{focus}' in request.form:
                focus_areas.append(focus)
        
        if focus_areas:
            options['focus_areas'] = focus_areas
        
        try:
            # Process the document (store in memory, no need to save to disk)
            results, paragraphs = process_document(file.stream, options)
            
            if 'error' in results:
                flash(results['error'])
                return redirect(url_for('index'))
            
            # Save results to a temp file
            result_id = str(uuid.uuid4())
            result_path = os.path.join(app.config['UPLOAD_FOLDER'], f'{result_id}.json')
            with open(result_path, 'w', encoding='utf-8') as f:
                json.dump(results, f)
            
            session['result_id'] = result_id
            return redirect(url_for('results'))
        except Exception as e:
            logger.error(f"Error during document processing: {e}")
            flash('An error occurred while processing the document.')
            return redirect(url_for('index'))
    
    flash('Invalid file type. Please upload a .docx file.')
    return redirect(url_for('index'))

@app.route('/results')
def results():
    """Display the analysis results"""
    result_id = session.get('result_id')
    if not result_id:
        flash('No results to display. Please analyze a document first.')
        return redirect(url_for('index'))
    
    try:
        # Validate the result ID format
        if not result_id or not isinstance(result_id, str) or len(result_id) != 36:
            flash('Invalid result ID format.')
            return redirect(url_for('index'))
            
        result_path = os.path.join(app.config['UPLOAD_FOLDER'], f'{result_id}.json')
        if not os.path.exists(result_path):
            flash('Results file not found.')
            return redirect(url_for('index'))
            
        with open(result_path, 'r', encoding='utf-8') as f:
            results = json.load(f)
            
        return render_template('results.html', results=results)
    except Exception as e:
        logger.error(f"Error loading results: {e}")
        flash('An error occurred while loading the results.')
        return redirect(url_for('index'))

@app.route('/api/analyze', methods=['POST'])
def api_analyze():
    """API endpoint for document analysis"""
    # Check if a file was uploaded
    if 'document' not in request.files:
        return jsonify({"error": "No file provided"}), 400
    
    file = request.files['document']
    
    if file.filename == '':
        return jsonify({"error": "No file selected"}), 400
    
    if file and allowed_file(file.filename):
        try:
            # Get options from form data
            options = {
                'model': request.form.get('model', 'gemma3:1b'),
                'api_url': request.form.get('api_url', 'http://localhost:11434'),
                'chunk_size': request.form.get('chunk_size', '3'),
                'genre': request.form.get('genre', ''),
                'sample_only': request.form.get('sample_only', 'false').lower() == 'true'
            }
            
            # Process the document
            results, paragraphs = process_document(file.stream, options)
            
            if 'error' in results:
                return jsonify(results), 400
            
            return jsonify(results)
        except Exception as e:
            logger.error(f"API error: {e}")
            return jsonify({"error": f"An error occurred: {str(e)}"}), 500
    
    return jsonify({"error": "Invalid file type"}), 400

@app.errorhandler(413)
def request_entity_too_large(error):
    """Handle file too large error"""
    flash('File too large. Maximum size is 50MB.')
    return redirect(url_for('index')), 413

@app.errorhandler(500)
def internal_server_error(error):
    """Handle internal server errors"""
    logger.error(f"Internal server error: {error}")
    flash('An internal server error occurred. Please try again later.')
    return redirect(url_for('index')), 500

if __name__ == "__main__":
    app.run(debug=True)